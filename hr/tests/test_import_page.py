from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from hr.models import Employee, Office, Role
from hr.tests.test_flow import PASSWORD, make_employee
from hr.tests.test_import import ROWS, workbook


def upload(path):
    with open(path, "rb") as f:
        return SimpleUploadedFile("staff.xlsx", f.read())


class ImportPage(TestCase):
    def setUp(self):
        self.admin = make_employee("A1", office=Office.UK, role=Role.ADMIN)
        self.client.login(username="uA1", password=PASSWORD)

    def test_preview_then_confirm(self):
        resp = self.client.post(reverse("staff_import"), {"file": upload(workbook(ROWS))})
        self.assertContains(resp, "3 to add")
        self.assertContains(resp, "£35,000/yr")
        self.assertEqual(Employee.objects.count(), 1)
        resp = self.client.post(reverse("staff_import"), {"action": "confirm"})
        self.assertRedirects(resp, reverse("employee_list"))
        self.assertEqual(Employee.objects.count(), 4)
        self.assertEqual(Employee.objects.get(employee_code="2002").team_lead.employee_code, "2001")

    def test_gaps_are_warnings_not_blockers(self):
        gappy = ROWS + [{"employee_id": "9", "full_name": "X", "office": "India", "designation": "X"}]
        resp = self.client.post(reverse("staff_import"), {"file": upload(workbook(gappy))})
        self.assertContains(resp, "no joining date")
        self.assertContains(resp, "Confirm and add 4 staff")
        self.client.post(reverse("staff_import"), {"action": "confirm"})
        self.assertTrue(Employee.objects.filter(employee_code="9").exists())
        # The staff page and dashboard still work with the missing date.
        self.assertEqual(self.client.get(reverse("employee_detail", args=[Employee.objects.get(employee_code="9").pk])).status_code, 200)
        self.assertContains(self.client.get(reverse("dashboard")), "Joining date missing")

    def test_rejects_non_xlsx(self):
        resp = self.client.post(reverse("staff_import"), {"file": SimpleUploadedFile("x.csv", b"a,b")})
        self.assertContains(resp, "Upload an Excel .xlsx file")

    def test_template_download(self):
        resp = self.client.get(reverse("staff_import") + "?template=1")
        sheet = load_workbook(BytesIO(resp.content)).active
        self.assertEqual(sheet.cell(row=1, column=1).value, "employee_id")

    def test_non_admin_blocked(self):
        make_employee("E1")
        self.client.login(username="uE1", password=PASSWORD)
        self.assertEqual(self.client.get(reverse("staff_import")).status_code, 403)
