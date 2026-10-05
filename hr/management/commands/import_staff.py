"""Import staff from an onboarding spreadsheet (.xlsx).

    python manage.py import_staff "imports/India staff onboarding.xlsx"           # preview only
    python manage.py import_staff "imports/India staff onboarding.xlsx" --commit  # save

The first sheet must have a header row using the column names in COLUMNS.
Nothing is saved unless every row is valid. Existing employee IDs are skipped.
"""

from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.core.exceptions import ValidationError
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from hr.models import Department, Employee, EmploymentRecord, Office, Role, audit

COLUMNS = [
    "employee_id", "full_name", "office", "designation", "department", "date_joined", "salary",
    "pf_enrolled", "esic_enrolled", "uan", "pan", "esic_ip_number", "aadhaar_last4",
    "email", "phone", "business_unit", "team_lead_id", "probation_end_date", "opening_leave_balance",
    "weekly_hours", "role",
]
REQUIRED = ["employee_id", "full_name", "office", "designation", "date_joined"]
OFFICES = {"india": Office.INDIA, "in": Office.INDIA, "uk": Office.UK, "united kingdom": Office.UK}
ROLES = {"": Role.EMPLOYEE, "employee": Role.EMPLOYEE, "team lead": Role.TEAM_LEAD, "team_lead": Role.TEAM_LEAD,
         "admin": Role.ADMIN}
YES = {"y", "yes", "true", "1"}
NO = {"", "n", "no", "false", "0"}


def text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def parse_date(value, field):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = text(value)
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y", "%d %b %Y", "%d %B %Y"):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"{field}: '{raw}' is not a date (use DD/MM/YYYY)")


def parse_decimal(value, field):
    raw = text(value).replace(",", "").replace("₹", "").replace("£", "")
    if raw == "":
        return None
    try:
        return Decimal(raw)
    except InvalidOperation:
        raise ValueError(f"{field}: '{raw}' is not a number")


def parse_bool(value, field):
    raw = text(value).lower()
    if raw in YES:
        return True
    if raw in NO:
        return False
    raise ValueError(f"{field}: use Yes or No")


class Command(BaseCommand):
    help = "Import staff from an onboarding .xlsx (preview by default; --commit to save)."

    def add_arguments(self, parser):
        parser.add_argument("path")
        parser.add_argument("--commit", action="store_true", help="Save the rows. Without this, only preview.")

    def handle(self, path, commit, **options):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise CommandError("openpyxl is needed: pip install openpyxl")
        try:
            sheet = load_workbook(path, data_only=True).worksheets[0]
        except FileNotFoundError:
            raise CommandError(f"File not found: {path}")

        rows = list(sheet.iter_rows(values_only=True))
        header = [text(h).lower() for h in rows[0]]
        missing = [c for c in COLUMNS if c not in header]
        if missing:
            raise CommandError(f"Missing columns: {', '.join(missing)}")
        records = []
        for number, values in enumerate(rows[1:], start=2):
            row = dict(zip(header, values))
            if not any(text(v) for v in row.values()):
                continue
            if text(row.get("employee_id")).lower() == "employee id":  # friendly heading row
                continue
            records.append((number, row))

        parsed, errors = [], []
        for number, row in records:
            try:
                parsed.append((number, self.parse_row(row)))
            except ValueError as exc:
                errors.append(f"Row {number}: {exc}")

        codes = [p["employee_code"] for _, p in parsed]
        for code in {c for c in codes if codes.count(c) > 1}:
            errors.append(f"Employee ID {code} appears more than once.")
        existing = set(Employee.objects.filter(employee_code__in=codes).values_list("employee_code", flat=True))
        known_leads = set(codes) | set(Employee.objects.values_list("employee_code", flat=True))
        for number, p in parsed:
            if p["team_lead_code"] and p["team_lead_code"] not in known_leads:
                errors.append(f"Row {number}: team_lead_id {p['team_lead_code']} is not a known employee ID.")

        self.print_preview(parsed, existing)
        if errors:
            self.stdout.write(self.style.ERROR("\nFix these before importing:"))
            for err in errors:
                self.stdout.write(f"  - {err}")
            raise CommandError(f"{len(errors)} problem(s) found. Nothing was saved.")

        to_create = [(n, p) for n, p in parsed if p["employee_code"] not in existing]
        if not commit:
            self.stdout.write(self.style.WARNING(
                f"\nPreview only: {len(to_create)} to add, {len(existing)} already in the system (skipped). "
                "Run again with --commit to save."
            ))
            return
        try:
            with transaction.atomic():
                created = self.create(to_create)
        except ValidationError as exc:
            raise CommandError(f"Nothing was saved: {exc}")
        self.stdout.write(self.style.SUCCESS(f"\nAdded {len(created)} staff. Skipped {len(existing)} existing."))

    def parse_row(self, row):
        for field in REQUIRED:
            if not text(row.get(field)):
                raise ValueError(f"{field} is required")
        office = OFFICES.get(text(row["office"]).lower())
        if office is None:
            raise ValueError("office must be India or UK")
        role = ROLES.get(text(row.get("role")).lower())
        if role is None:
            raise ValueError("role must be Employee, Team lead or Admin")
        salary = parse_decimal(row.get("salary"), "salary")
        if office == Office.INDIA and salary is None:
            raise ValueError("salary (monthly Basic in INR) is required for India staff")
        aadhaar = text(row.get("aadhaar_last4"))
        if aadhaar and not (aadhaar.isdigit() and len(aadhaar) == 4):
            raise ValueError("aadhaar_last4 must be exactly the last 4 digits")
        india = office == Office.INDIA
        return {
            "employee_code": text(row["employee_id"]),
            "full_name": text(row["full_name"]),
            "office": office,
            "role": role,
            "designation": text(row["designation"]),
            "department": text(row.get("department")),
            "date_joined": parse_date(row["date_joined"], "date_joined"),
            "probation_end_date": parse_date(row.get("probation_end_date"), "probation_end_date"),
            "salary": salary,
            "weekly_hours": parse_decimal(row.get("weekly_hours"), "weekly_hours"),
            "email": text(row.get("email")),
            "phone": text(row.get("phone")),
            "business_unit": text(row.get("business_unit")),
            "team_lead_code": text(row.get("team_lead_id")),
            "pf_enrolled": parse_bool(row.get("pf_enrolled"), "pf_enrolled") if india else False,
            "esic_enrolled": parse_bool(row.get("esic_enrolled"), "esic_enrolled") if india else False,
            "uan": text(row.get("uan")) if india else "",
            "pan": text(row.get("pan")).upper() if india else "",
            "esic_ip_number": text(row.get("esic_ip_number")) if india else "",
            "aadhaar_last4": aadhaar if india else "",
            "opening_leave_balance": (parse_decimal(row.get("opening_leave_balance"), "opening_leave_balance")
                                      or Decimal("0")) if india else Decimal("0"),
        }

    def print_preview(self, parsed, existing):
        self.stdout.write(f"{'ID':<8}{'Name':<30}{'Office':<7}{'Designation':<34}{'Joined':<12}"
                          f"{'Salary':>12}  PF  ESIC  Status")
        for _, p in parsed:
            currency = "INR" if p["office"] == Office.INDIA else "GBP"
            salary = f"{currency} {p['salary']:,.0f}" if p["salary"] is not None else "-"
            status = "exists, skip" if p["employee_code"] in existing else "new"
            self.stdout.write(
                f"{p['employee_code']:<8}{p['full_name'][:29]:<30}{p['office']:<7}{p['designation'][:33]:<34}"
                f"{p['date_joined']:%d/%m/%Y}  {salary:>12}  {'Y' if p['pf_enrolled'] else '-':<3} "
                f"{'Y' if p['esic_enrolled'] else '-':<5} {status}"
            )

    def create(self, rows):
        created = {}
        for _, p in rows:
            department = None
            if p["department"]:
                department, _ = Department.objects.get_or_create(name=p["department"], office=p["office"])
            e = Employee(
                employee_code=p["employee_code"], full_name=p["full_name"], office=p["office"], role=p["role"],
                designation=p["designation"], department=department, date_joined=p["date_joined"],
                probation_end_date=p["probation_end_date"], email=p["email"], phone=p["phone"],
                business_unit=p["business_unit"], pf_enrolled=p["pf_enrolled"], esic_enrolled=p["esic_enrolled"],
                uan=p["uan"], pan=p["pan"], esic_ip_number=p["esic_ip_number"], aadhaar_last4=p["aadhaar_last4"],
                opening_leave_balance=p["opening_leave_balance"],
            )
            e.full_clean(exclude=["team_lead", "user"])
            e.save()
            record = EmploymentRecord(
                employee=e, effective_from=p["date_joined"], designation=p["designation"], department=department,
                weekly_hours=p["weekly_hours"],
                basic_monthly=p["salary"] if p["office"] == Office.INDIA else None,
                annual_salary_gbp=p["salary"] if p["office"] == Office.UK else None,
                reason="Imported at go-live (salary as at import; earlier history not recorded)",
            )
            record.full_clean(exclude=["created_by"])
            record.save()
            audit(None, "employee.imported", e)
            created[e.employee_code] = (e, p["team_lead_code"])
        for e, lead_code in created.values():
            if lead_code:
                e.team_lead = Employee.objects.get(employee_code=lead_code)
                e.save(update_fields=["team_lead"])
        return created
